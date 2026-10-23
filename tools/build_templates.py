#!/usr/bin/env python3
"""Costruisce i template docxtemplater di Toscana Diagnostica · Business a partire dai format aziendali
(tools/assets) e il JSON delle liste di autovalutazione (dagli XLSX compilati per Calenzano).

Uso: python3 tools/build_templates.py
Output: server/src/templates/*.docx e server/src/data/autovalutazione.json

I format restano la fonte: per aggiornare un testo contrattuale si sostituisce il DOCX in tools/assets
e si rilancia lo script. I segnaposto {campo} vengono compilati dal server (server/src/lib/docgen.js).
"""
import copy, json, os, re, sys
import docx
import openpyxl
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, 'assets')
OUT = os.path.join(HERE, '..', 'server', 'src', 'templates')
DATA = os.path.join(HERE, '..', 'server', 'src', 'data')
os.makedirs(OUT, exist_ok=True); os.makedirs(DATA, exist_ok=True)

# ---------------------------------------------------------------- utilità

def set_text(p, text):
    """Sostituisce il testo del paragrafo con un solo run, conservando la formattazione del primo run."""
    runs = p.runs
    first = runs[0]._element if runs else None
    for el in list(p._p):
        if el.tag.endswith('}pPr') or el is first:
            continue
        p._p.remove(el)
    if runs:
        runs[0].text = text
    else:
        p.add_run(text)

def iter_paragraphs(doc):
    """Tutti i paragrafi: corpo, tabelle (anche annidate), intestazioni e piè di pagina."""
    def from_table(t):
        for row in t.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for p in cell.paragraphs:
                    yield p
                for nt in cell.tables:
                    yield from from_table(nt)
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        yield from from_table(t)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            try:
                for p in part.paragraphs:
                    yield p
                for t in part.tables:
                    yield from from_table(t)
            except Exception:
                pass

def replace_all(doc, mapping):
    """Sostituzione di stringhe esatte in tutto il documento (anche quando il testo è spezzato su più run)."""
    n = 0
    for p in iter_paragraphs(doc):
        t = p.text
        if not t:
            continue
        new = t
        for k, v in mapping.items():
            if k in new:
                new = new.replace(k, v)
        if new != t:
            set_text(p, new)
            n += 1
    return n

def cell(doc, t, r, c, text, pi=0):
    set_text(doc.tables[t].rows[r].cells[c].paragraphs[pi], text)

def save(doc, name):
    path = os.path.join(OUT, name)
    apply_letterhead(doc)
    doc.save(path)
    print('template', name)

LETTERHEAD = docx.Document(os.path.join(ASSETS, 'Carta_Intestata_TD.docx'))

def apply_letterhead(doc):
    """Intestazione e piè di pagina della carta intestata ufficiale su ogni sezione del documento.
    Le immagini dell'intestazione sono richiamate per rId: si verifica che nel documento di destinazione
    gli stessi rId puntino alle stesse immagini (tutti i format derivano dalla carta intestata)."""
    import hashlib
    lh = LETTERHEAD.sections[0]
    def img_map(part):
        return {rid: hashlib.md5(rel.target_part.blob).hexdigest() for rid, rel in part.rels.items() if 'image' in rel.reltype}
    src_imgs = img_map(lh.header.part)
    for s in doc.sections:
        s.different_first_page_header_footer = False
        for src, dst in ((lh.header, s.header), (lh.footer, s.footer)):
            dst.is_linked_to_previous = False
            if dst is s.header:
                dst_imgs = img_map(dst.part)
                for rid, h in src_imgs.items():
                    if dst_imgs.get(rid) != h:
                        raise RuntimeError(f'intestazione: rId {rid} non corrisponde alla carta intestata')
            el = dst._element
            for child in list(el):
                el.remove(child)
            for child in src._element:
                el.append(copy.deepcopy(child))

def shrink_images(path, max_w=1600, quality=80):
    """Riduce le immagini del pacchetto DOCX (organigrammi a piena risoluzione) per tenere leggeri i documenti generati."""
    import io, zipfile, shutil
    from PIL import Image
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith('word/media/') and item.filename.lower().endswith(('.jpg', '.jpeg', '.png')) and len(data) > 300_000:
                im = Image.open(io.BytesIO(data))
                if im.width > max_w:
                    im = im.resize((max_w, round(im.height * max_w / im.width)), Image.LANCZOS)
                buf = io.BytesIO()
                if item.filename.lower().endswith('.png'):
                    im.save(buf, 'PNG', optimize=True)
                else:
                    im.convert('RGB').save(buf, 'JPEG', quality=quality, optimize=True)
                data = buf.getvalue()
            zout.writestr(item, data)
    shutil.move(tmp, path)

def strip_drawings(doc):
    """Rimuove le immagini dal corpo (es. firma autografa scansionata presente nel format compilato)."""
    n = 0
    for p in doc.paragraphs:
        for el in list(p._p.iter()):
            if el.tag.endswith('}drawing') or el.tag.endswith('}pict'):
                el.getparent().remove(el); n += 1
    return n

def add_trace(doc):
    p = doc.add_paragraph()
    r = p.add_run('{trace_line}')
    r.font.size = Pt(7); r.italic = True

# Dati personali presenti nei format (Calenzano): nel template diventano segnaposto riempiti dalle Impostazioni
LR_BLOCK_OLD = 'EPIFANI FRANCESCO, nato a POGGIARDO (LE) il 03/08/1984, residente a SESTO FIORENTINO (FI), VIA ANTONIO GRAMSCI, 128, codice fiscale PFNFNC84M03G751B, recapito telefonico +39 3470584620'
LR_BLOCK_NEW = '{lr_name_upper}, nato a {lr_birth_place} il {lr_birth_date}, residente a {lr_residence}, codice fiscale {lr_tax_code}, recapito telefonico {lr_phone}'
CO_BLOCK_OLD = 'TOSCANA DIAGNOSTICA S.R.L., con sede legale in FIRENZE (FI), VIALE BELFIORE, 55, codice fiscale/partita IVA 07528130482, Tel. 055776511, Cell. +39 3470584620, Fax 055776511, PEC toscanadiagnostica@pec.it, e-mail: francesco.epifani@toscanadiagnostica.it'
CO_BLOCK_NEW = '{company_name_upper}, con sede legale in {company_legal_address_upper}, codice fiscale/partita IVA {company_vat}, Tel. {company_phone}, Cell. {lr_phone}, Fax {company_fax}, PEC {company_pec}, e-mail: {lr_email}'
COMMON = {
    LR_BLOCK_OLD: LR_BLOCK_NEW,
    CO_BLOCK_OLD: CO_BLOCK_NEW,
    'Via di Prato, 26, 50041, Calenzano (FI)': '{site_street}, {site_zip}, {site_city} ({site_prov})',
    'Via di Prato 26, Calenzano (FI)': '{site_street}, {site_city} ({site_prov})',
    'Punto Prelievi Calenzano': 'Punto Prelievi {site_city}',
    'MAESTRALE SRL': '{host_name_upper}',
    'Iacopo Cupelli': '{rspp_name}', 'IACOPO CUPELLI': '{rspp_name_upper}',
    'Sara Arrighetti': '{rls_name}', 'SARA ARRIGHETTI': '{rls_name_upper}',
    'Vincenzo Cupelli': '{mc_name}', 'VINCENZO CUPELLI': '{mc_name_upper}',
    'prof. Gian Luigi Taddei': '{ds_title_name}', 'Prof. Gian Luigi Taddei': '{ds_title_name}',
}

# ---------------------------------------------------------------- 1. Contratto Punto Prelievi
d = docx.Document(os.path.join(ASSETS, 'Contratto_Punto_Prelievi_Format.docx'))
P = d.paragraphs
set_text(P[2], '{partner_name}, con sede in {partner_city} ({partner_prov}) – {partner_street}, P.I. {partner_vat} e Codice Fiscale {partner_tax_code}, in persona del Sig. {partner_signer}, munito degli occorrenti poteri (nel seguito la “Società”)')
set_text(P[4], '{company_name_upper}, con sede legale in {company_legal_address}, Partita IVA {company_vat} e Codice Fiscale {company_tax_code}, in persona del suo legale rappresentante pro tempore {lr_title_name}, munito degli occorrenti poteri (nel seguito “Toscana Diagnostica”)')
set_text(P[8], 'La Società gestisce un poliambulatorio specialistico, sito in {site_city} ({site_prov}), {site_street} (di seguito “Poliambulatorio”) ove svolgono la propria attività medici specialisti, medici in libera professione, infermieri e/o personale ASA, le cui prestazioni ambulatoriali sono erogate in regime privatistico e/o tramite speciali convenzioni;')
set_text(P[32], 'a mantenere l’esclusiva nel comune di {site_city} su eventuali ulteriori aperture.')
# Servizio infermieristico e accettazione: a carico di Toscana Diagnostica o della struttura ospitante (scelta contrattuale).
# I paragrafi sono avvolti in condizioni docxtemplater: con paragraphLoop il paragrafo vuoto viene rimosso.
def clone_after(par, text):
    el = copy.deepcopy(par._p)
    par._p.addnext(el)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(el, par._parent)
    set_text(np_, text)
    return np_
set_text(P[31], '{#td_nursing}a garantire il servizio infermieristico per l’esecuzione dei prelievi tramite personale qualificato e dotato degli idonei titoli previsti dalla legge, nonché a garantire il pieno rispetto delle procedure Toscana Diagnostica e a fornire al proprio personale tutti i dispositivi di protezione individuale richiesti dal tipo di attività svolta (quali, a mero titolo esemplificativo e non esaustivo, camice, guanti e mascherina etc.);{/td_nursing}')
clone_after(P[31], '{#td_reception}a garantire, tramite proprio personale e propri sistemi, il servizio di segreteria, accettazione, prenotazione, gestione delle prescrizioni mediche, gestione ed amministrazione dei rapporti con gli utenti (ivi inclusa la gestione delle attività relative al trattamento dei dati personali ed alla acquisizione dei consensi privacy e sanitari, ove previsto dalla normativa), stampa e consegna dei referti agli utenti, contabilità e fatturazione all’utenza e gestione degli incassi;{/td_reception}')
set_text(P[37], '{#host_reception}a garantire il servizio di segreteria, accettazione, servizi di prenotazione, gestione delle prescrizioni mediche, gestione ed amministrazione dei rapporti con gli utenti (ivi inclusa la gestione delle attività relative al trattamento dei dati personali ed alla acquisizione dei consensi privacy e sanitari, ove previsto dalla normativa), stampa e consegna dei referti agli utenti, contabilità e fatturazione all’utenza e gestione degli incassi. Si precisa che la fatturazione sarà effettuata tramite sistemi Toscana Diagnostica e le fatture saranno emesse a nome di quest’ultima;{/host_reception}')
clone_after(P[37], '{#host_nursing}a garantire il servizio infermieristico per l’esecuzione dei prelievi tramite proprio personale qualificato e dotato degli idonei titoli previsti dalla legge, nel pieno rispetto delle procedure Toscana Diagnostica, fornendo al proprio personale tutti i dispositivi di protezione individuale richiesti dal tipo di attività svolta; il personale infermieristico impiegato resta in rapporto esclusivo con la Società, senza alcun vincolo con Toscana Diagnostica;{/host_nursing}')
set_text(P[52], 'Il Contratto ha durata di anni {duration_years} ({duration_years_words}), con decorrenza dal {start_date} (o in ogni data utile precedente), data di apertura del Punto Prelievi.')
# il format rimanda all'Allegato 2 per i locali: i locali sono identificati nell'Allegato 1
set_text(P[55], 'In caso di esercizio del diritto di recesso dal Contratto da parte della Società, quest’ultima, per una durata minima di 9 (nove) mesi dalla data di cessazione del rapporto, non potrà concedere i locali di cui all’Allegato 1 a soggetti terzi, per lo svolgimento di attività analoghe a quelle svolte da Toscana Diagnostica.')
set_text(P[58], '{fee_clause} Detto corrispettivo costituisce una remunerazione complessiva ed onnicomprensiva per la prestazione dei Servizi da parte della Società.')
set_text(P[94], 'quanto alla Società: {partner_name}, {partner_city} ({partner_prov}), {partner_street}, P.I. {partner_vat} e Codice Fiscale {partner_tax_code}, in persona del Sig. {partner_signer}{partner_pec_clause};')
set_text(P[95], 'quanto a Toscana Diagnostica: {company_name_upper}, con sede legale in {company_legal_address}, in persona del suo legale rappresentante pro tempore {lr_title_name};')
cell(d, 0, 0, 0, '{partner_name}'); cell(d, 0, 1, 0, '{partner_signer_role}'); cell(d, 0, 2, 0, '{partner_signer}')
cell(d, 0, 2, 1, '{lr_name}')
set_text(P[108], '{duvri_costs_clause}')
# riga di protocollo e data in testa
head = P[0].insert_paragraph_before('')
r = head.add_run('Proposta {protocol} del {proposal_date}'); r.font.size = Pt(9); r.italic = True
add_trace(d)
save(d, 'contratto_punto_prelievi.docx')

# ---------------------------------------------------------------- 2. Offerta Service (MOD230)
d = docx.Document(os.path.join(ASSETS, 'MOD230_Contratto_Service.docx'))
P = d.paragraphs
set_text(P[0], 'Spett.le\n{client_name}\n{client_street}\n{client_zip} {client_city} ({client_prov})\nP.IVA {client_vat} · C.F. {client_tax_code}\nc.a. {client_signer}\n(di seguito anche "Cliente" o "Società")')
P[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
set_text(P[7], '{site_name}')
set_text(P[8], '{site_address}')
for i in (26, 31, 172):
    set_text(P[i], '{company_city} ({company_prov}), {proposal_date}')
set_text(P[60], 'indicare il nominativo del Direttore Sanitario della struttura: {client_ds};')
set_text(P[184], 'Listino esami in service con sconto riservato al Cliente. {price_list_clause}')
cell(d, 1, 4, 1, '{client_pec}'); cell(d, 1, 5, 1, '{client_sdi}')
cell(d, 1, 6, 1, '{vat_exemption_clause}')
cell(d, 1, 7, 1, '{payment_terms}')
cell(d, 1, 8, 1, '{client_bank}')
for t in (2, 4, 5):
    cell(d, t, 1, 0, '{client_signer}'); cell(d, t, 1, 1, '{lr_role}'); cell(d, t, 2, 1, '{lr_name}')
# Allegato 3: ritiro
cell(d, 7, 1, 1, '{pickup_site}'); cell(d, 7, 2, 1, '{pickup_street}'); cell(d, 7, 3, 1, '{pickup_city}')
cell(d, 7, 4, 1, '{pickup_contact}'); cell(d, 7, 5, 1, '{pickup_contact_ref}')
for i in range(1, 7):
    cell(d, 8, i, 1, '{d%d_fixed} Fisso  {d%d_oncall} A chiamata' % (i, i)); cell(d, 8, i, 2, '{d%d_time}' % i)
# Allegato 6: soluzione
set_text(P[214], '{sol1}  Soluzione 1 — Interfacciamento bidirezionale via LIS TD')
set_text(P[215], '{sol2}  Soluzione 2 — Interfacciamento manuale via portale web')
set_text(P[216], '{sol3}  Soluzione 3 — Schede cartacee predefinite')
# Allegato 8: privacy
set_text(P[224], '{priv_owner}  TITOLARE DEL TRATTAMENTO (Cliente con Pazienti come propri diretti clienti)')
set_text(P[225], '{priv_processor}  RESPONSABILE DEL TRATTAMENTO (Cliente fornitore di service per conto di un Committente)')
cell(d, 11, 1, 2, '{priv_principal}'); cell(d, 11, 2, 1, '{priv_measures_no}  NO  {priv_measures_yes}  SI (elencare)'); cell(d, 11, 2, 2, '{priv_measures}')
cell(d, 11, 3, 2, '{priv_dpo_name}'); cell(d, 11, 4, 2, '{priv_dpo_contact}')
head = P[0].insert_paragraph_before('')
r = head.add_run('Proposta {protocol} del {proposal_date}'); r.font.size = Pt(9); r.italic = True
add_trace(d)
save(d, 'offerta_service.docx')

# ---------------------------------------------------------------- 3. Offerta Medicina del Lavoro (MOD231)
d = docx.Document(os.path.join(ASSETS, 'MOD231_Contratto_MdL.docx'))
P = d.paragraphs
set_text(P[0], 'Spett.le\n{client_name}\n{client_street}\n{client_zip} {client_city} ({client_prov})\nP.IVA {client_vat} · C.F. {client_tax_code}\nc.a. {client_signer}\n(di seguito anche "Cliente" o "Società")')
P[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
for i in (21, 26, 142):
    set_text(P[i], '{company_city} ({company_prov}), {proposal_date}')
set_text(P[53], 'Il Cliente autorizza Toscana Diagnostica ad anticipare al Medico del Lavoro preventivamente indicato, in caso di necessità e su richiesta, i risultati dei referti dei Lavoratori dipendenti (e ad eccezione dei referti contenenti dati idonei a rivelare l’identità genetica e relativi al Test HIV), a mezzo e-mail al seguente indirizzo {doctor_email}, con le modalità e le opportune cautele atte a garantire al massimo il rispetto della riservatezza dei dati sanitari ivi contenuti. In particolare, tenendo in considerazione le indicazioni del Garante di cui al paragrafo 6 delle Linee guida in tema di referti on-line adottate dal Garante in data 25 giugno 2009, (i) il referto sarà spedito in forma di allegato a un messaggio e-mail e non come testo compreso nella body part del messaggio; (ii) il file contenente il referto sarà protetto con modalità idonee a impedire l\'illecita o fortuita acquisizione delle informazioni trasmesse da parte di soggetti diversi da quelli autorizzati, che dovranno consistere in una password per l\'apertura del file o in una chiave crittografica, rese note al Medico del Lavoro indicato dal Cliente tramite canali di comunicazione differenti da quelli utilizzati per la spedizione dei referti.')
set_text(P[146], 'Il catalogo è fornito al Cliente in formato Excel e PDF e aggiornato periodicamente con comunicazione via e-mail al Cliente. Per ogni esame sono indicati: codice, denominazione, settore, materiale richiesto, tempo di refertazione, eventuale invio a Terzi Incaricati. {price_list_clause}')
cell(d, 1, 1, 1, 'Si rinvia a quanto descritto nelle Condizioni Generali di Contratto. Il Cliente indica l’indirizzo e-mail di riferimento, utile alla conferma degli appuntamenti per l’esecuzione degli esami suddetti: {service_email} e il contatto di riferimento sig. {service_contact}, Tel. {service_phone}')
cell(d, 1, 2, 1, 'Mensile. Il Cliente autorizza Toscana Diagnostica ad inviare le fatture per via elettronica al seguente indirizzo di posta elettronica: {invoice_email} (ai sensi del D. Lgs. n. 52/2004). Resta inteso che, ove le fatture o gli allegati contengano dati personali relativi ai pazienti, essa sarà inviata a mezzo PEC con le modalità di cui all’ultimo comma dell’Art. 2 delle Condizioni Generali di Contratto.')
cell(d, 1, 4, 1, '{client_pec}'); cell(d, 1, 5, 1, '{client_sdi}'); cell(d, 1, 6, 1, '{vat_exemption_clause}'); cell(d, 1, 7, 1, '{payment_terms}'); cell(d, 1, 8, 1, '{client_bank}')
for t in (2, 4, 6):
    cell(d, t, 1, 0, '{client_signer}'); cell(d, t, 1, 1, '{lr_role}'); cell(d, t, 2, 1, '{lr_name}')
set_text(P[154], '{priv_owner}  TITOLARE DEL TRATTAMENTO (Cliente con Pazienti come propri diretti clienti)')
set_text(P[155], '{priv_processor}  RESPONSABILE DEL TRATTAMENTO (Cliente fornitore di service per conto di un Committente)')
cell(d, 8, 1, 2, '{priv_principal}'); cell(d, 8, 2, 1, '{priv_measures_no}  NO  {priv_measures_yes}  SI (elencare)'); cell(d, 8, 2, 2, '{priv_measures}')
cell(d, 8, 3, 2, '{priv_dpo_name}'); cell(d, 8, 4, 2, '{priv_dpo_contact}')
head = P[0].insert_paragraph_before('')
r = head.add_run('Proposta {protocol} del {proposal_date}'); r.font.size = Pt(9); r.italic = True
add_trace(d)
save(d, 'offerta_mdl.docx')

# ---------------------------------------------------------------- 4. Fascicolo autorizzazione: Modello 5
d = docx.Document(os.path.join(ASSETS, 'Modello_5.docx'))
P = d.paragraphs
set_text(P[1], 'del Comune di {site_city}')
set_text(P[6], 'ai sensi della L.R. 5 agosto 2009 n. 51, il rilascio dell’autorizzazione all’apertura di un punto di prelievo ubicato in {site_city_upper}, {site_street_upper}, Tel. {company_phone}, Cell. {lr_phone}, Fax {company_fax}, PEC {company_pec}, e-mail: {lr_email}.')
replace_all(d, {**COMMON, '22/05/2026': '{doc_date}', 'Firenze (FI), ': '{company_city} ({company_prov}), ',
    'SESTO FIORENTINO (FI), VIA DI PRATIGNONE, 13/4': '{lab_address_upper}'})
save(d, 'modello_5.docx')

# ---------------------------------------------------------------- 5. Modello 6 (Direttore Sanitario)
d = docx.Document(os.path.join(ASSETS, 'Modello_6.docx'))
P = d.paragraphs
set_text(P[8], 'Il sottoscritto {ds_name_upper} nato a {ds_birth_place} il {ds_birth_date}, residente in {ds_residence}, Codice Fiscale {ds_tax_code}, Telefono {ds_phone}, indirizzo mail {ds_email}, assume l’incarico di Direttore Sanitario della Struttura Sanitaria denominata ' + CO_BLOCK_NEW + ', con impegno orario settimanale di n. {ds_weekly_hours} ore a far data da {ds_start_date}.')
set_text(P[14], '{ds_degree};')
set_text(P[15], '{ds_specialization};')
set_text(P[16], '{ds_registration}.')
replace_all(d, {'22/05/2026': '{doc_date}', 'Firenze (FI), ': '{company_city} ({company_prov}), '})
print('modello 6: rimosse', strip_drawings(d), 'immagini (firma autografa)')
save(d, 'modello_6.docx')

# ---------------------------------------------------------------- 6. Relazione attività non insudicianti
d = docx.Document(os.path.join(ASSETS, 'Relazione_non_insudicianti.docx'))
P = d.paragraphs
set_text(P[4], 'Ragione Sociale ditta: {company_name}')
set_text(P[5], 'Titolare / Legale Rappresentante: {lr_name}')
set_text(P[6], 'Sede operativa: {site_street}, {site_city} ({site_prov})')
set_text(P[33], '{company_city}, {doc_date}')
set_text(P[36], '{lr_title_name}')
save(d, 'relazione_non_insudicianti.docx')

# ---------------------------------------------------------------- 7. Collegamento funzionale
d = docx.Document(os.path.join(ASSETS, 'Collegamento_Funzionale.docx'))
P = d.paragraphs
# il format riportava per errore il sottotitolo e la premessa della dichiarazione sporco/pulito
set_text(P[1], 'Dichiarazione di collegamento funzionale con il Laboratorio di Analisi.')
set_text(P[3], 'Il sottoscritto ' + LR_BLOCK_NEW + ', in qualità di LEGALE RAPPRESENTANTE di ' + CO_BLOCK_NEW + ', ai fini della apertura del Punto Prelievi decentrato ubicato in {site_city} ({site_prov}), {site_street},')
set_text(P[7], 'che il Punto Prelievi in parola è funzionalmente collegato al Laboratorio di Analisi di {company_name}, ubicato in {lab_address}.')
replace_all(d, {'22/05/2026': '{doc_date}', 'Firenze (FI), ': '{company_city} ({company_prov}), ', 'Francesco Epifani': '{lr_name}'})
save(d, 'collegamento_funzionale.docx')

# ---------------------------------------------------------------- 8. Dichiarazione sporco/pulito
d = docx.Document(os.path.join(ASSETS, 'Dichiarazione_Sporco_Pulito.docx'))
P = d.paragraphs
set_text(P[3], 'Il sottoscritto ' + LR_BLOCK_NEW + ', in qualità di LEGALE RAPPRESENTANTE di ' + CO_BLOCK_NEW + ', d’intesa con il Responsabile del Servizio di Prevenzione e Protezione (RSPP), ai fini della apertura del Punto Prelievi ubicato in {site_city} ({site_prov}), {site_street}, e relativamente alla individuazione delle posizioni degli spazi destinati al deposito di materiale “sporco” e “pulito”, di cui al requisito BS10, BS11 della lista di Autovalutazione sez. B')
set_text(P[7], '{dirty_clean_clause}')
replace_all(d, {'22/05/2026': '{doc_date}', 'Firenze (FI), ': '{company_city} ({company_prov}), ', 'Francesco Epifani': '{lr_name}'})
save(d, 'dichiarazione_sporco_pulito.docx')

# ---------------------------------------------------------------- 9. Atto di condivisione procedure
d = docx.Document(os.path.join(ASSETS, 'Atto_Condivisione_Procedure.docx'))
P = d.paragraphs
set_text(P[1], '{host_name}')
set_text(P[2], 'c/o Punto Prelievi {site_city}')
set_text(P[3], '{host_street}')
set_text(P[4], '{host_zip} {host_city} ({host_prov})')
set_text(P[6], 'Oggetto: Apertura punto di prelievo decentrato {company_name} c/o {host_name_upper} sito in {site_street}, {site_city} ({site_prov}) - Condivisione procedure Ambulatorio della {host_name_upper} / {company_name}.')
set_text(P[9], '{lr_title_name}, in qualità di {lr_role_it} di {company_name},')
set_text(P[10], '{host_signer}, Legale Rappresentante di {host_name_upper}.')
set_text(P[43], 'il {lr_title_name}, in qualità di {lr_role_it} di {company_name}')
set_text(P[48], 'Che le procedure operative e la modulistica di {company_name} sono disponibili presso il punto prelievi e consegnata al personale addetto all’accettazione e all’attività di prelievo.')
set_text(P[50], '{company_city}, {doc_date}')
cell(d, 0, 0, 1, '{company_name}'); cell(d, 0, 1, 0, '{host_name_upper}'); cell(d, 0, 1, 1, '{lr_role_it}'); cell(d, 0, 2, 0, '{host_signer_role}'); cell(d, 0, 2, 1, '{lr_title_name}')
replace_all(d, {'Toscana Diagnostica Srl': '{company_name}'})
save(d, 'atto_condivisione_procedure.docx')

# ---------------------------------------------------------------- 10. DUVRI
d = docx.Document(os.path.join(ASSETS, 'DUVRI.docx'))
P = d.paragraphs
set_text(P[28], 'LUOGO: {host_name_upper}, {site_street}, {site_city} ({site_prov})')
set_text(P[88], 'PUNTO PRELIEVI – {host_name_upper} ({site_floor}), {site_street}, {site_city} ({site_prov}). Il Punto Prelievi di {company_name} ')
set_text(P[115], '{contract_period} rinnovabile tacitamente')
replace_all(d, {**COMMON, '23/06/2026': '{doc_date}', '30/05/2026': '{doc_date}', 'TD-SGI-REG021_00': '{duvri_code}',
    'Toscana Diagnostica Srl': '{company_name}', 'FRANCESCO EPIFANI': '{lr_name_upper}', 'Francesco Epifani': '{lr_name}'})
save(d, 'duvri.docx')

# ---------------------------------------------------------------- 11. DVR
d = docx.Document(os.path.join(ASSETS, 'DVR.docx'))
P = d.paragraphs
set_text(P[6], '{site_street}, {site_city} ({site_prov})')
set_text(P[33], 'Il centro prelievi di TOSCANA DIAGNOSTICA è di piccole dimensioni ({site_workers} addetti); la sede è nel comune di {site_city}. L’insediamento è situato all’interno di una struttura sanitaria gestita da altro soggetto, con ottime possibilità di esodo e individuazione del punto di raccolta. ')
cell(d, 5, 1, 1, '{site_city}'); cell(d, 5, 2, 1, '{site_prov_name}'); cell(d, 5, 3, 1, '{site_street}'); cell(d, 5, 4, 1, '{site_building_type}')
replace_all(d, {**COMMON, '30/05/2026': '{doc_date}', 'TD-SGI-REG020_00': '{dvr_code}', 'Francesco Epifani': '{lr_name}', '€ 10.000,00 i.v.': '{company_capital}'})
save(d, 'dvr.docx')

# ---------------------------------------------------------------- 12. Relazione operativa (con orari, personale, strumentazione)
d = docx.Document(os.path.join(ASSETS, 'Relazione_Operativa.docx'))
P = d.paragraphs
set_text(P[34], 'Apertura di un nuovo Punto Prelievi decentrato in {site_city} ({site_prov}), {site_street}, collegato funzionalmente con il Laboratorio di Analisi di {company_name} situato presso {lab_address}. {site_notes}')
# T1 orari: righe lunedì..domenica
DAYS = ['Lunedì', 'Martedì', 'Mercoledì', 'Giovedì', 'Venerdì', 'Sabato', 'Domenica']
t = d.tables[1]
for i in range(7):
    set_text(t.rows[i + 1].cells[1].paragraphs[0], '{h%d_draw}' % (i + 1))
    set_text(t.rows[i + 1].cells[2].paragraphs[0], '{h%d_pickup}' % (i + 1))
# T3 personale: una riga ciclo
t = d.tables[3]
while len(t.rows) > 2:
    t._tbl.remove(t.rows[2]._tr)
row = t.rows[1]
vals = ['{#staff}{surname}', '{name}', '{affiliation}', '{contract}', '{hours}', '{role}', '{degree}{/staff}']
for c, v in zip(row.cells, vals):
    set_text(c.paragraphs[0], v)
# T5 strumentazione
t = d.tables[5]
while len(t.rows) > 2:
    t._tbl.remove(t.rows[2]._tr)
row = t.rows[1]
vals = ['{#equipment}{n}', '{type}', '{maker}', '{model}', '{serial}', '{power}', '{ownership}', '{condition}', '{tested_at}{/equipment}']
for c, v in zip(row.cells, vals):
    set_text(c.paragraphs[0], v)
replace_all(d, {**COMMON, 'Toscana Diagnostica S.r.l.': '{company_name}', 'Francesco Epifani': '{lr_name}'})
save(d, 'relazione_operativa.docx')
shrink_images(os.path.join(OUT, 'relazione_operativa.docx'))

# Elenco personale e strumentazione come documenti autonomi: stesso format, solo la sezione interessata
def section_only(src, keep_from, keep_to, name):
    doc = docx.Document(src)
    body = doc.element.body
    paras = doc.paragraphs
    start = paras[keep_from]._p; end = paras[keep_to]._p
    keeping = False
    for el in list(body):
        if el.tag.endswith('}sectPr'):
            continue
        if el is start:
            keeping = True
        if not keeping:
            body.remove(el)
        if el is end:
            keeping = False
            # tutto ciò che segue viene rimosso dal ciclo
    # rimuove ciò che segue 'end'
    after = False
    for el in list(body):
        if el.tag.endswith('}sectPr'):
            continue
        if after:
            body.remove(el)
        if el is end:
            after = True
    return doc

src = os.path.join(OUT, 'relazione_operativa.docx')
dd = section_only(src, 123, 142, 'personale')
save(dd, 'elenco_personale.docx'); shrink_images(os.path.join(OUT, 'elenco_personale.docx'))
dd = section_only(src, 143, len(docx.Document(src).paragraphs) - 1, 'strumentazione')
save(dd, 'elenco_strumentazione.docx'); shrink_images(os.path.join(OUT, 'elenco_strumentazione.docx'))

# ---------------------------------------------------------------- 13. Liste di autovalutazione -> JSON
def read_xlsx(path, list_code):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    title = str(ws.cell(1, 1).value or '').strip()
    section = None
    for r in ws.iter_rows(min_row=3, values_only=True):
        vals = list(r[:6])
        code, text, applic, answer, document, notes = [str(v).strip() if v is not None else '' for v in vals]
        if not any([code, text, applic, answer, document, notes]):
            continue
        # righe di sezione: solo la prima colonna valorizzata, in maiuscolo
        if code and not text and not answer and not document and code.upper() == code and len(code) > 3:
            section = code
            continue
        rows.append({'code': code, 'text': text, 'applicability': applic, 'answer': answer, 'document': document, 'notes': notes, 'section': section})
    return {'code': list_code, 'title': title, 'items': rows}

lists = [
    read_xlsx(os.path.join(ASSETS, 'Autovalutazione_A.xlsx'), 'A'),
    read_xlsx(os.path.join(ASSETS, 'Autovalutazione_B.xlsx'), 'B'),
    read_xlsx(os.path.join(ASSETS, 'Autovalutazione_B1.1.xlsx'), 'B1.1'),
]
# le risposte compilate per Calenzano diventano il modello: i riferimenti alla sede vengono resi generici
def generic(s):
    return (s.replace('Punto Prelievi di Calenzano', 'Punto Prelievi').replace('Punto Prelievi Calenzano', 'Punto Prelievi')
             .replace('di Calenzano', '').replace('Calenzano', 'della sede'))
for L in lists:
    for it in L['items']:
        for k in ('answer', 'document', 'notes'):
            it[k] = generic(it[k])
with open(os.path.join(DATA, 'autovalutazione.json'), 'w', encoding='utf-8') as f:
    json.dump(lists, f, ensure_ascii=False, indent=1)
print('autovalutazione.json:', [(L['code'], len(L['items'])) for L in lists])
